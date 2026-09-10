"""
ingest_releases.py
==================
Ingests and normalizes multi-state release schedules from Excel files in _Input/:
  - AK 2026 Release Calendar.xlsx (Alaska MMIS)
  - ND Release Calendar - 2026.xlsx (North Dakota MMIS)
  - NH MMIS_ ReleaseSchedule_Till_01-31-2027_Approved_10162025 (1).xlsx (New Hampshire MMIS)

Populates `release_schedules` and `release_milestones` tables in SQLite.
"""

from __future__ import annotations

import json
import pathlib
import re
import sqlite3
from datetime import date, datetime, timedelta
from typing import Any

import openpyxl

from db import get_connection, replace_release_milestones, upsert_release_schedule


def _clean_date_str(val: Any) -> str | None:
    """Normalize date objects or strings into YYYY-MM-DD format."""
    if val is None:
        return None
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if not s or s.upper() in ("NA", "NONE", "N/A", "DAILY"):
        return None
    # Try YYYY-MM-DD
    match = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if match:
        y, m, d = match.groups()
        return f"{int(y):04d}-{int(m):02d}-{int(d):02d}"
    # Try M/D/YYYY or M/D/YY
    match = re.search(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", s)
    if match:
        m, d, y = match.groups()
        if len(y) == 2:
            y = f"20{y}"
        return f"{int(y):04d}-{int(m):02d}-{int(d):02d}"
    return None


def _get_quarter(dt_str: str | None) -> tuple[int, str]:
    """Return year and quarter (e.g. 2026, 'Q1') from date string."""
    if not dt_str:
        return 2026, "Q1"
    try:
        dt = datetime.strptime(dt_str[:10], "%Y-%m-%d")
        q = f"Q{(dt.month - 1) // 3 + 1}"
        return dt.year, q
    except Exception:
        return 2026, "Q1"


# ==============================================================================
# 1. Alaska (AK) Parser
# ==============================================================================
def parse_ak_calendar(file_path: pathlib.Path) -> list[dict]:
    """Parses Alaska 2026 Release Calendar Excel."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    releases = []
    cur_rel: dict[str, Any] | None = None

    for r in rows[1:]:
        name = r[0]
        if not name:
            continue
        s_name = str(name).strip()
        if s_name.startswith("AK."):
            cur_rel = {
                "state": "AK",
                "release_id": s_name,
                "release_name": f"Alaska MMIS {s_name}",
                "state_rm_name": "Alaska State RM Lead",
                "state_rm_email": "ak-rm-lead@ets.state.gov",
                "start_date": _clean_date_str(r[2]),
                "finish_date": _clean_date_str(r[3]),
                "duration": str(r[1]) if r[1] is not None else "",
                "milestones": [],
            }
            releases.append(cur_rel)
        elif cur_rel:
            cur_rel["milestones"].append({
                "name": s_name,
                "task_name": s_name,
                "duration": str(r[1]) if r[1] is not None else "",
                "start": _clean_date_str(r[2]),
                "finish": _clean_date_str(r[3]),
                "predecessors": str(r[4]) if len(r) > 4 and r[4] is not None else "",
                "rework": str(r[5]) if len(r) > 5 and r[5] is not None else "",
                "holiday": str(r[6]) if len(r) > 6 and r[6] is not None else "",
            })

    # Consolidate key stage dates for each AK release
    results = []
    for rel in releases:
        m_list = rel["milestones"]
        prod_date = None
        dev_start = None
        dev_end = None
        sit_start = None
        sit_end = None
        uat_start = None
        uat_end = None
        scope_freeze = None
        go_nogo = None

        for m in m_list:
            task = m["name"].lower()
            if "deployment" in task and "post" not in task:
                prod_date = m["finish"] or m["start"]
            elif "development" in task:
                dev_start = m["start"]
                dev_end = m["finish"]
            elif "qa" in task or "functional & regression" in task:
                sit_start = m["start"]
                sit_end = m["finish"]
            elif "state uat" in task or "uat" in task:
                uat_start = m["start"]
                uat_end = m["finish"]
            elif "analysis" in task or "planning" in task or "hld" in task:
                if not scope_freeze:
                    scope_freeze = m["finish"]
            elif "sign-off" in task or "go" in task:
                go_nogo = m["start"]

        if not prod_date:
            prod_date = rel["finish_date"]

        year, quarter = _get_quarter(prod_date)

        # Categorize milestones
        normalized_milestones = []
        for m in m_list:
            t = m["name"].lower()
            if "deploy" in t:
                cat = "Production"
                env = "PROD"
            elif "uat" in t:
                cat = "UAT"
                env = "UAT"
            elif "qa" in t or "regression" in t:
                cat = "SIT"
                env = "SIT"
            elif "dev" in t:
                cat = "Development"
                env = "DEV"
            elif "review" in t or "sign-off" in t or "approval" in t:
                cat = "Gate"
                env = "SOA/FAS"
            else:
                cat = "Planning"
                env = "MGMT"

            now_dt = datetime.now().date()
            now_str = now_dt.strftime("%Y-%m-%d")
            m_status = "Passed" if (m["finish"] and m["finish"] <= now_str) else ("Active" if (m["start"] and m["start"] <= now_str) else "Scheduled")
            normalized_milestones.append({
                "task_name": m["name"],
                "phase_category": cat,
                "env_target": env,
                "duration_str": m["duration"],
                "start_date": m["start"],
                "finish_date": m["finish"],
                "predecessors": m["predecessors"],
                "holiday_impact": m["holiday"],
                "status": m_status,
            })

        # Calculate status & readiness
        now_dt = datetime.now().date()
        now_str = now_dt.strftime("%Y-%m-%d")
        if prod_date and prod_date < now_str:
            status = "Completed"
            readiness = 100.0
            risk = "Low"
        elif prod_date and prod_date == now_str:
            status = "In Progress"
            readiness = 96.0
            risk = "High"
        elif prod_date and prod_date <= (now_dt + timedelta(days=35)).strftime("%Y-%m-%d"):
            status = "In Progress"
            readiness = 88.5
            risk = "Medium"
        else:
            status = "Scheduled"
            readiness = 72.0
            risk = "Low"

        results.append({
            "state": "AK",
            "release_id": rel["release_id"],
            "release_name": rel["release_name"],
            "state_rm_name": rel["state_rm_name"],
            "state_rm_email": rel["state_rm_email"],
            "year": year,
            "quarter": quarter,
            "scope_freeze_date": scope_freeze,
            "dev_start_date": dev_start,
            "dev_end_date": dev_end,
            "sit_start_date": sit_start,
            "sit_end_date": sit_end,
            "uat_start_date": uat_start,
            "uat_end_date": uat_end,
            "go_nogo_date": go_nogo,
            "prod_deploy_date": prod_date,
            "status": status,
            "risk_level": risk,
            "milestones_total": len(normalized_milestones),
            "milestones_completed": sum(1 for x in normalized_milestones if x["status"] == "Passed"),
            "readiness_pct": readiness,
            "notes": f"Alaska SOA approval gates & FAS testing matrix cadence ({len(normalized_milestones)} tasks)",
            "raw_json": json.dumps({"milestones": normalized_milestones}),
            "_milestone_list": normalized_milestones,
        })

    return results


# ==============================================================================
# 2. North Dakota (ND) Parser
# ==============================================================================
def parse_nd_calendar(file_path: pathlib.Path) -> list[dict]:
    """Parses North Dakota Release Calendar 2026 Excel."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    results = []
    for r in rows[1:]:
        rel_id = r[0]
        if not rel_id or not str(rel_id).startswith("ND."):
            continue
        rel_str = str(rel_id).strip()

        # Check date string across row for active window
        row_dates_str = " ".join([str(x) for x in r if x is not None])
        if not any(yr in row_dates_str for yr in ("2025-11", "2025-12", "2026", "2027")):
            continue

        prod_date = _clean_date_str(r[13]) if len(r) > 13 else None
        if not prod_date:
            continue

        scope_freeze = _clean_date_str(r[1])
        dev_start = _clean_date_str(r[2])
        dev_code_checkin = _clean_date_str(r[3])
        code_cutoff = _clean_date_str(r[4])
        deploy_76 = _clean_date_str(r[5])
        sit_cutoff = _clean_date_str(r[6])
        sit_reg_start = _clean_date_str(r[7])
        sit_reg_complete = _clean_date_str(r[8])
        uat_deploy = _clean_date_str(r[9])
        uat_bf = _clean_date_str(r[10]) if len(r) > 10 else None
        uat_end = _clean_date_str(r[11]) if len(r) > 11 else None
        go_nogo = _clean_date_str(r[12]) if len(r) > 12 else None

        year, quarter = _get_quarter(prod_date)

        # Build granular milestone list
        ms_defs = [
            ("Defects & CR Scope Freeze", "Planning", "MMIS", scope_freeze, scope_freeze),
            ("Development Coding Start", "Development", "DEV", dev_start, dev_code_checkin or dev_start),
            ("Development Code Check-In", "Development", "DEV", dev_code_checkin, dev_code_checkin),
            ("Code Cut-Off (Week 2)", "Development", "DEV", code_cutoff, code_cutoff),
            ("Build & Deployment to 76", "Development", "Build 76", deploy_76, deploy_76),
            ("SIT Deployment Cutoff", "SIT", "SIT", sit_cutoff, sit_cutoff),
            ("SIT Regression Testing Start (Deployments Allowed)", "SIT", "SIT", sit_reg_start, sit_reg_start),
            ("SIT Regression Testing Complete (Code Freeze)", "SIT", "SIT", sit_reg_complete, sit_reg_complete),
            ("State UAT Deployment (Week 1)", "UAT", "UAT", uat_deploy, uat_deploy),
            ("UAT Break-Fix (BF) Cycle", "UAT", "UAT", uat_bf, uat_bf),
            ("State UAT Testing Sign-Off", "UAT", "UAT", uat_end, uat_end),
            ("Formal Go / No-Go Decision Gate", "Gate", "GOV", go_nogo, go_nogo),
            ("Production Live Deployment (Week 2 Cutover)", "Production", "PROD", prod_date, prod_date),
        ]

        now_dt = datetime.now().date()
        now_str = now_dt.strftime("%Y-%m-%d")

        normalized_milestones = []
        for name, cat, env, s_dt, f_dt in ms_defs:
            if s_dt or f_dt:
                m_status = "Passed" if (f_dt and f_dt <= now_str) else ("Active" if (s_dt and s_dt <= now_str) else "Scheduled")
                normalized_milestones.append({
                    "task_name": name,
                    "phase_category": cat,
                    "env_target": env,
                    "duration_str": "Gate" if s_dt == f_dt else "Sprint",
                    "start_date": s_dt,
                    "finish_date": f_dt or s_dt,
                    "predecessors": "",
                    "holiday_impact": "",
                    "status": m_status,
                })

        if prod_date < now_str:
            status = "Completed"
            readiness = 100.0
            risk = "Low"
        elif prod_date == now_str:
            status = "In Progress"
            readiness = 96.0
            risk = "High"
        elif prod_date <= (now_dt + timedelta(days=35)).strftime("%Y-%m-%d"):
            status = "In Progress"
            readiness = 90.0
            risk = "Medium"
        else:
            status = "Scheduled"
            readiness = 75.0
            risk = "Low"

        results.append({
            "state": "ND",
            "release_id": rel_str,
            "release_name": f"North Dakota MMIS {rel_str}",
            "state_rm_name": "North Dakota State RM Lead",
            "state_rm_email": "nd-rm-lead@ets.state.gov",
            "year": year,
            "quarter": quarter,
            "scope_freeze_date": scope_freeze,
            "dev_start_date": dev_start,
            "dev_end_date": code_cutoff or dev_code_checkin,
            "sit_start_date": sit_reg_start,
            "sit_end_date": sit_reg_complete,
            "uat_start_date": uat_deploy,
            "uat_end_date": uat_end,
            "go_nogo_date": go_nogo,
            "prod_deploy_date": prod_date,
            "status": status,
            "risk_level": risk,
            "milestones_total": len(normalized_milestones),
            "milestones_completed": sum(1 for x in normalized_milestones if x["status"] == "Passed"),
            "readiness_pct": readiness,
            "notes": f"North Dakota Build-76 pipeline & formal Go/No-Go gate ({len(normalized_milestones)} gates)",
            "raw_json": json.dumps({"milestones": normalized_milestones}),
            "_milestone_list": normalized_milestones,
        })

    return results


# ==============================================================================
# 3. New Hampshire (NH) Parser
# ==============================================================================
def parse_nh_calendar(file_path: pathlib.Path) -> list[dict]:
    """Parses New Hampshire MMIS Approved Release Schedule (Sheet 20251016)."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet_name = "20251016" if "20251016" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))

    nh_raw_releases: list[dict[str, Any]] = []
    cur_nh_rel: dict[str, Any] | None = None

    for r in rows[2:]:
        val0 = r[0]
        env = r[1]
        phase = r[2]

        if val0 is not None and str(val0).strip() and str(val0).strip() != "Release":
            cur_nh_rel = {
                "release_id": f"NH.{str(val0).strip()}",
                "version": str(val0).strip(),
                "tracks": [],
            }
            nh_raw_releases.append(cur_nh_rel)

        if cur_nh_rel and env:
            # Find non-empty date values across the row
            non_empty_dates = [c for c in r[3:] if c is not None]
            cur_nh_rel["tracks"].append({
                "env": str(env).strip(),
                "phase": str(phase).strip() if phase else "",
                "dates": non_empty_dates,
            })

    results = []
    for rel in nh_raw_releases:
        prod_date = None
        dev_start, dev_end = None, None
        sit_start, sit_end = None, None
        uat_start, uat_end = None, None
        scope_freeze = None
        normalized_milestones = []

        for trk in rel["tracks"]:
            env = trk["env"]
            dates = trk["dates"]
            phase_name = trk["phase"]

            # Check if this track is Prod
            if "ENV05" in env or "PROD" in env.upper():
                for d in dates:
                    clean_d = _clean_date_str(d)
                    if clean_d:
                        prod_date = clean_d
                        break
                normalized_milestones.append({
                    "task_name": f"Production Deployment ({env})",
                    "phase_category": "Production",
                    "env_target": "ENV05 (Prod)",
                    "duration_str": "Weekend Cutover",
                    "start_date": prod_date,
                    "finish_date": prod_date,
                    "predecessors": "UAT Sign-off",
                    "holiday_impact": "",
                    "status": "Passed" if (prod_date and prod_date < "2026-03-01") else "Active",
                })
            elif "ENV52" in env or "DEV" in phase_name.upper():
                date_label = str(dates[0]) if dates else ""
                normalized_milestones.append({
                    "task_name": f"MMIS Development ({env})",
                    "phase_category": "Development",
                    "env_target": "ENV52 (Dev)",
                    "duration_str": date_label,
                    "start_date": None,
                    "finish_date": None,
                    "predecessors": "Scope Freeze",
                    "holiday_impact": "",
                    "status": "Active",
                })
            elif "ENV57" in env or "SIT" in phase_name.upper():
                date_label = str(dates[0]) if dates else ""
                normalized_milestones.append({
                    "task_name": f"MMIS SIT Testing ({env})",
                    "phase_category": "SIT",
                    "env_target": "ENV57 (SIT)",
                    "duration_str": date_label,
                    "start_date": None,
                    "finish_date": None,
                    "predecessors": "Dev Drop",
                    "holiday_impact": "",
                    "status": "Active",
                })
            elif "ENV53" in env or "REGRESSION" in phase_name.upper():
                date_label = str(dates[0]) if dates else ""
                normalized_milestones.append({
                    "task_name": f"Regression Execution ({env})",
                    "phase_category": "SIT",
                    "env_target": "ENV53 (Regression)",
                    "duration_str": date_label,
                    "start_date": None,
                    "finish_date": None,
                    "predecessors": "SIT Complete",
                    "holiday_impact": "",
                    "status": "Active",
                })
            elif "ENV04" in env or "UAT" in phase_name.upper():
                date_label = str(dates[0]) if dates else ""
                normalized_milestones.append({
                    "task_name": f"State UAT Acceptance ({env})",
                    "phase_category": "UAT",
                    "env_target": "ENV04 (UAT)",
                    "duration_str": date_label,
                    "start_date": None,
                    "finish_date": None,
                    "predecessors": "Regression Complete",
                    "holiday_impact": "",
                    "status": "Active",
                })
            elif "SYSDOC" in env.upper():
                normalized_milestones.append({
                    "task_name": "System Documentation & NTT Data QA Review",
                    "phase_category": "Gate",
                    "env_target": "SysDoc",
                    "duration_str": "Check-out / Check-in",
                    "start_date": None,
                    "finish_date": None,
                    "predecessors": "",
                    "holiday_impact": "",
                    "status": "Active",
                })

        if not prod_date:
            continue

        year, quarter = _get_quarter(prod_date)

        now_dt = datetime.now().date()
        now_str = now_dt.strftime("%Y-%m-%d")
        if prod_date < now_str:
            status = "Completed"
            readiness = 100.0
            risk = "Low"
        elif prod_date == now_str:
            status = "In Progress"
            readiness = 96.0
            risk = "High"
        elif prod_date <= (now_dt + timedelta(days=35)).strftime("%Y-%m-%d"):
            status = "In Progress"
            readiness = 94.0
            risk = "Medium"
        else:
            status = "Scheduled"
            readiness = 75.0
            risk = "Low"

        results.append({
            "state": "NH",
            "release_id": rel["release_id"],
            "release_name": f"New Hampshire MMIS Release {rel['version']}",
            "state_rm_name": "New Hampshire State RM Lead",
            "state_rm_email": "nh-rm-lead@ets.state.gov",
            "year": year,
            "quarter": quarter,
            "scope_freeze_date": scope_freeze,
            "dev_start_date": dev_start,
            "dev_end_date": dev_end,
            "sit_start_date": sit_start,
            "sit_end_date": sit_end,
            "uat_start_date": uat_start,
            "uat_end_date": uat_end,
            "go_nogo_date": None,
            "prod_deploy_date": prod_date,
            "status": status,
            "risk_level": risk,
            "milestones_total": len(normalized_milestones),
            "milestones_completed": sum(1 for x in normalized_milestones if x["status"] == "Passed"),
            "readiness_pct": readiness,
            "notes": f"New Hampshire multi-environment pipeline (ENV52/57/53/04/05) & NTT Data documentation governance",
            "raw_json": json.dumps({"milestones": normalized_milestones}),
            "_milestone_list": normalized_milestones,
        })

    return results


# ==============================================================================
# Ingestion Orchestrator
# ==============================================================================
def run_release_ingest(input_dir: pathlib.Path | str, db_path: str) -> dict[str, Any]:
    """Ingest all Excel files from input_dir and insert into SQLite database."""
    in_dir = pathlib.Path(input_dir)
    conn = get_connection(db_path)

    ak_file = in_dir / "AK 2026 Release Calendar.xlsx"
    nd_file = in_dir / "ND Release Calendar - 2026.xlsx"
    nh_file = in_dir / "NH MMIS_ ReleaseSchedule_Till_01-31-2027_Approved_10162025 (1).xlsx"

    all_releases = []
    if ak_file.exists():
        all_releases.extend(parse_ak_calendar(ak_file))
    if nd_file.exists():
        all_releases.extend(parse_nd_calendar(nd_file))
    if nh_file.exists():
        all_releases.extend(parse_nh_calendar(nh_file))

    total_milestones = 0
    for rel in all_releases:
        m_list = rel.pop("_milestone_list", [])
        upsert_release_schedule(conn, rel)
        replace_release_milestones(conn, rel["release_id"], rel["state"], m_list)
        total_milestones += len(m_list)

    conn.close()

    return {
        "total_releases_ingested": len(all_releases),
        "total_milestones_ingested": total_milestones,
        "states_ingested": sorted(list(set(r["state"] for r in all_releases))),
    }


if __name__ == "__main__":
    ROOT = pathlib.Path(__file__).resolve().parent.parent
    db_file = str(ROOT / "data" / "expiry.db")
    input_folder = ROOT / "_Input"
    res = run_release_ingest(input_folder, db_file)
    print("Ingestion result:", res)
